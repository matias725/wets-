"""Convierte el Excel de SAP (maestro de unidades + líneas de OT) en los datos
que usa WEST IA Web: public/data/west-real.json

Uso:
    python scripts/convertir_sap.py "C:/ruta/SAP COMPLETO.xlsx"
    python scripts/convertir_sap.py archivo.xlsx --catalogo "C:/.../west_ia_manager/preventive_catalog.py"

El JSON generado contiene datos reales de la empresa (clientes, RUT, costos):
está excluido de git y NO debe subirse a un repositorio público.
Requiere: pip install openpyxl
"""
from __future__ import annotations

import argparse
import ast
import json
import math
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

DEFAULT_CATALOG = Path.home() / "OneDrive/Desktop/West IA Personal/WEST_IA_V3_5/west_ia_manager/preventive_catalog.py"
OUT = Path(__file__).resolve().parent.parent / "public" / "data" / "west-real.json"

# Coordenadas aproximadas por sucursal / faena (lat, lng, zona)
PLACES = {
    "taller central": ("Taller Central (La Serena)", -29.9045, -71.2489, "centro"),
    "la serena": ("La Serena", -29.9045, -71.2489, "centro"),
    "san geronimo": ("San Gerónimo", -29.93, -71.07, "centro"),
    "teck": ("Teck (Andacollo)", -30.2333, -71.0833, "centro"),
    "andacollo": ("Andacollo", -30.2333, -71.0833, "centro"),
    "hmc": ("HMC (Punitaqui)", -30.8333, -71.2667, "centro"),
    "punitaqui": ("Punitaqui", -30.8333, -71.2667, "centro"),
    "salamanca": ("Salamanca", -31.7797, -70.9636, "centro"),
    "vallenar": ("Vallenar", -28.5708, -70.7581, "norte"),
    "copiapo": ("Copiapó", -27.3668, -70.3322, "norte"),
    "calama": ("Calama", -22.4567, -68.9237, "norte"),
    "antofagasta": ("Antofagasta", -23.6509, -70.3975, "norte"),
    "san pedro atacama": ("San Pedro de Atacama", -22.9087, -68.1997, "norte"),
    "iquique": ("Iquique", -20.2141, -70.1524, "norte"),
    "arica": ("Arica", -18.4783, -70.3126, "norte"),
    "santiago": ("Santiago", -33.4489, -70.6693, "centro"),
    "concepcion": ("Concepción", -36.827, -73.0503, "sur"),
    "los angeles": ("Los Ángeles", -37.4693, -72.3527, "sur"),
    "temuco": ("Temuco", -38.7359, -72.5904, "sur"),
    "pucon": ("Pucón", -39.2819, -71.9544, "sur"),
    "puerto montt": ("Puerto Montt", -41.4693, -72.9424, "sur"),
}

ACTIVE = {"no iniciada", "proceso"}


def norm(v) -> str:
    s = unicodedata.normalize("NFD", str(v or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s.replace("_", " ").replace("-", " ")).strip()


def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def num(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").replace(".", "").replace(",", ".").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def iso(v) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
    m = re.match(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", str(v))
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(v))
    return m.group(0) if m else ""


def plate_of(v) -> str:
    raw = re.sub(r"[^A-Z0-9]", "", str(v or "").upper())
    if re.fullmatch(r"[A-Z]{4}\d{2}", raw):
        return f"{raw[:4]}-{raw[4:]}"
    if re.fullmatch(r"[A-Z]{2}\d{4}", raw):
        return f"{raw[:2]}-{raw[2:]}"
    if re.fullmatch(r"[A-Z]{3}\d{3}", raw):
        return f"{raw[:3]}-{raw[3:]}"
    return raw


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", norm(s)).strip("-") or "sin-sucursal"


def load_catalog(path: Path) -> set[str]:
    if not path.exists():
        print(f"AVISO: no se encontró el catálogo preventivo en {path}; se usará una regla por palabras clave.")
        return set()
    tree = ast.parse(path.read_text(encoding="utf-8"))
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(getattr(t, "id", "") == "PREVENTIVE_ITEMS" for t in node.targets):
            return {str(code).strip().upper() for code, *_ in ast.literal_eval(node.value)}
    return set()


def find_header(ws, required: list[str], max_rows: int = 30):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True)):
        heads = [norm(c) for c in row]
        if all(any(h == r for h in heads) for r in required):
            return i + 1, {h: j for j, h in enumerate(heads) if h}
    return None, None


def category_of(desc: str, model: str) -> str:
    d, m = norm(desc), norm(model)
    if "camioneta 4x4 " in d + " " and "-" in str(desc):
        return "pickup-mining"
    if "4x4" in d or ("camioneta" in d and "4x4" in m):
        return "pickup-4x4"
    if "camioneta" in d:
        return "pickup-4x2"
    if any(k in d for k in ("minibus", "bus", "furgon", "vans", "taxibus")):
        return "bus"
    if "camion" in d:
        return "truck"
    if "station" in d or "jeep" in d:
        return "suv"
    if "automovil" in d:
        return "sedan"
    return "other"


def classify(lines: list[dict]) -> str:
    groups = {norm(l["group"]) for l in lines}
    descs = " ".join(norm(l["description"]) for l in lines)
    if "s seguro vehiculo" in groups or "deducible" in descs or "siniestro" in descs:
        return "Compañía de seguros"
    if "s desabolladura y pintura" in groups:
        return "DYP"
    substantive = [l for l in lines if norm(l["code"]) != "ingresotaller"]
    if not substantive:
        return "Otros"
    if all(norm(l["group"]) == "s revision tecnica" for l in substantive):
        return "Revisión técnica"
    if all("lavado" in norm(l["group"]) for l in substantive):
        return "Lavado"
    prev = any(l["prev"] for l in substantive)
    corr = any(
        not l["prev"] and norm(l["group"]) in ("a repuestos", "a neumaticos", "a baterias", "s mano de obra terceros")
        for l in substantive
    )
    if prev and corr:
        return "Preventiva + Correctiva"
    if prev:
        return "Preventiva"
    if corr:
        return "Correctiva"
    if all(norm(l["group"]) in ("a accesorios", "a control de unidades", "a epp") for l in substantive):
        return "Equipamiento unidades nuevas"
    return "Otros"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel")
    ap.add_argument("--catalogo", default=str(DEFAULT_CATALOG))
    ap.add_argument("--salida", default=str(OUT))
    args = ap.parse_args()

    src = Path(args.excel)
    catalog = load_catalog(Path(args.catalogo))
    print(f"Leyendo {src.name} …")
    wb = load_workbook(src, read_only=True, data_only=True)

    # ------------------------------------------------------------ líneas de OT
    ot_ws = hdr = None
    for ws in wb.worksheets:
        row, h = find_header(ws, ["no ot", "patente", "costo total"])
        if row:
            ot_ws, ot_row, hdr = ws, row, h
            break
    if not ot_ws:
        sys.exit("No se encontró la hoja de OT (columnas No OT, Patente, Costo Total).")
    col = lambda r, name: r[hdr[name]] if name in hdr and hdr[name] < len(r) else None

    orders: dict[str, dict] = {}
    for r in ot_ws.iter_rows(min_row=ot_row + 1, values_only=True):
        if not r or col(r, "no ot") in (None, ""):
            continue
        ot = text(col(r, "no ot"))
        o = orders.get(ot)
        if o is None:
            sucursal = text(col(r, "sucursal"))
            o = orders[ot] = {
                "workOrder": ot,
                "plate": plate_of(col(r, "patente")),
                "branchRaw": sucursal,
                "clientName": text(col(r, "alias cliente")) or text(col(r, "cliente")),
                "clientFull": text(col(r, "cliente")),
                "area": text(col(r, "situacion")) or "RAC",
                "costCenter": text(col(r, "c.costo")),
                "receivedDate": iso(col(r, "fec. recepcion")) or iso(col(r, "fec. contab.")),
                "closedDate": iso(col(r, "fec. cierre")),
                "sapStatus": text(col(r, "estado ot")),
                "mileage": int(num(col(r, "kilometraje"))),
                "reason": text(col(r, "comentarios")),
                "createdBy": text(col(r, "genera ot")),
                "brandStyle": text(col(r, "marca/estilo")),
                "lines": [],
            }
        code = text(col(r, "codigo"))
        desc = text(col(r, "descripcion articulo/serv.")) or text(col(r, "descripcion articulo/serv")) or text(col(r, "descripcion"))
        if not desc:
            desc = next((text(v) for k, v in zip(hdr, r) if k.startswith("descripci")), "")
        o["lines"].append({
            "code": code,
            "description": desc,
            "qty": num(col(r, "cantidad")),
            "unitCost": round(num(col(r, "costo unitario"))),
            "total": round(num(col(r, "costo total"))),
            "group": text(col(r, "grupo")),
            "prev": code.upper() in catalog if catalog else bool(re.search(r"filtro|aceite|mantenc|pauta", norm(desc))),
        })

    # ------------------------------------------------------------- sucursales
    branch_ids: dict[str, str] = {}
    branches = []
    for raw, _ in Counter(o["branchRaw"] for o in orders.values() if o["branchRaw"]).most_common():
        name, lat, lng, zone = PLACES.get(norm(raw), (raw, None, None, "centro"))
        bid = slug(raw)
        branch_ids[raw] = bid
        branches.append({"id": bid, "name": name, "city": name, "lat": lat, "lng": lng, "zone": zone, "hq": norm(raw) == "taller central"})

    work_orders = []
    responsibles = defaultdict(Counter)
    for o in orders.values():
        o["branchId"] = branch_ids.get(o["branchRaw"])
        o["totalCost"] = sum(l["total"] for l in o["lines"])
        o["interventionType"] = classify(o["lines"])
        o["recovery"] = "Por revisar" if o["interventionType"] in ("Compañía de seguros", "DYP") else "No recuperable"
        if not o["reason"]:
            o["reason"] = next((l["description"] for l in o["lines"] if norm(l["code"]) != "ingresotaller"), "Ingreso a taller")
        if o["createdBy"] and o["branchId"]:
            responsibles[o["branchId"]][o["createdBy"]] += 1
        # líneas compactas: [código, descripción, cant., unitario, total, preventivo]
        o["lines"] = [[l["code"], l["description"], l["qty"], l["unitCost"], l["total"], 1 if l["prev"] else 0] for l in o["lines"]]
        work_orders.append(o)
    work_orders.sort(key=lambda o: o["receivedDate"])

    # ---------------------------------------------------------------- maestro
    by_plate = defaultdict(list)
    for o in work_orders:
        by_plate[o["plate"]].append(o)

    m_ws = None
    for ws in wb.worksheets:
        row, h = find_header(ws, ["patente", "marca"])
        if row and ws is not ot_ws:
            m_ws, m_row, mh = ws, row, h
            break
    vehicles = []
    seen = set()
    if m_ws:
        mc = lambda r, name: r[mh[name]] if name in mh and mh[name] < len(r) else None
        for r in m_ws.iter_rows(min_row=m_row + 1, values_only=True):
            if not r or not mc(r, "patente"):
                continue
            plate = plate_of(mc(r, "patente"))
            if plate in seen:
                continue
            seen.add(plate)
            area = text(mc(r, "area negocio")).upper()
            ots = by_plate.get(plate, [])
            last = ots[-1] if ots else None
            active = [o for o in ots if norm(o["sapStatus"]) in ACTIVE]
            km = max((o["mileage"] for o in ots), default=0)
            prevs = [o for o in ots if o["interventionType"].startswith("Preventiva") and o["mileage"]]
            next_km = (prevs[-1]["mileage"] + 10_000) if prevs else (math.ceil((km + 1) / 10_000) * 10_000 if km else 10_000)
            if active:
                status = "workshop"
            elif area == "USADOS":
                status = "sold"
            elif area == "PERDIDA TOTAL":
                status = "out"
            else:
                status = "available"
            year = int(num(mc(r, "ano"))) if num(mc(r, "ano")) else None
            fuel = norm(mc(r, "combustible"))
            vehicles.append({
                "plate": plate,
                "vin": text(mc(r, "n.chasis")),
                "brand": text(mc(r, "marca")).title(),
                "model": text(mc(r, "modelo sap")),
                "kind": text(mc(r, "descripcion")),
                "category": category_of(text(mc(r, "descripcion")), text(mc(r, "modelo sap"))),
                "year": year if year and 1980 < year <= date.today().year + 1 else None,
                "transmission": "",
                "fuel": "Diésel" if "diesel" in fuel or "petroleo" in fuel else "Bencina" if "bencina" in fuel or "gasolina" in fuel else text(mc(r, "combustible")).title(),
                "color": text(mc(r, "color vehiculo")).title(),
                "branchId": last["branchId"] if last else None,
                "branchRaw": last["branchRaw"] if last else "",
                "status": status,
                "mileage": km,
                "nextMaintenanceKm": next_km,
                "clientId": None,
                "clientName": text(mc(r, "nombre sn")) or text(mc(r, "alias cc/cliente")),
                "area": "LOP" if area == "LOP" else "RAC" if area == "RAC" else area.title() or "RAC",
                "areaRaw": area,
                "owner": text(mc(r, "propiedad")),
                "costCenter": text(mc(r, "ultimo centro de costo")),
                "note": text(mc(r, "observacion")),
                "documents": [],
                "dailyRate": 0,
            })

    clients = sorted({v["clientName"] for v in vehicles if v["clientName"]} | {o["clientName"] for o in work_orders if o["clientName"]})
    dates = [o["receivedDate"] for o in work_orders if o["receivedDate"]]
    data = {
        "meta": {
            "source": "sap",
            "fileName": src.name,
            "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "vehicles": len(vehicles),
            "workOrders": len(work_orders),
            "lines": sum(len(o["lines"]) for o in work_orders),
            "from": min(dates) if dates else "",
            "to": max(dates) if dates else "",
            "platesOnlyInOT": sorted(set(by_plate) - seen),
        },
        "branches": branches,
        "vehicles": vehicles,
        "workOrders": work_orders,
        "responsibles": {b: [n for n, _ in c.most_common(6)] for b, c in responsibles.items()},
        "clients": clients,
    }
    out = Path(args.salida)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    st = Counter(v["status"] for v in vehicles)
    it = Counter(o["interventionType"] for o in work_orders)
    print(f"Listo: {out}  ({out.stat().st_size / 1e6:.1f} MB)")
    print(f"Vehículos: {len(vehicles)}  · estados: {dict(st)}")
    print(f"OT: {len(work_orders)}  · abiertas: {sum(1 for o in work_orders if norm(o['sapStatus']) in ACTIVE)}  · líneas: {data['meta']['lines']}")
    print(f"Tipos: {dict(it)}")
    print(f"Sucursales: {len(branches)} · sin coordenadas: {[b['name'] for b in branches if b['lat'] is None]}")
    print(f"Patentes con OT que no están en el maestro: {len(data['meta']['platesOnlyInOT'])}")


if __name__ == "__main__":
    main()
